from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "coindcx_inr_futures_bot_pitch.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(90, 96, 106)
GRAY_FILL = "F4F6F9"
LIGHT_BLUE_FILL = "E8EEF5"
BORDER = "C8D0DA"
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, *, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, *, size, color=None, bold=None):
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color=BORDER):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_para(doc, text="", *, style=None, size=None, bold=False, color=None, align=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_callout(doc, title: str, body: str, *, fill=GRAY_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [9360])
    set_table_borders(table, color="D3DAE4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    title_run = p.add_run(title)
    set_run_font(title_run, size=11, color=INK, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    body_run = p2.add_run(body)
    set_run_font(body_run, size=10.5, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_label_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [2100, 7260])
    set_table_borders(table)
    for idx, (label, value) in enumerate(rows):
        label_cell = table.cell(idx, 0)
        value_cell = table.cell(idx, 1)
        set_cell_shading(label_cell, LIGHT_BLUE_FILL)
        for cell in (label_cell, value_cell):
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
        label_run = label_cell.paragraphs[0].add_run(label)
        set_run_font(label_run, size=10, color=INK, bold=True)
        value_run = value_cell.paragraphs[0].add_run(value)
        set_run_font(value_run, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_matrix(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, widths)
    set_table_borders(table)
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, GRAY_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header)
        set_run_font(run, size=9.5, color=INK, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            set_run_font(run, size=9.3, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(doc, text: str):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, size=10.8, color=INK)


def add_number(doc, text: str):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    run = p.add_run(text)
    set_run_font(run, size=10.8, color=INK)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, size=11, color=INK)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    set_style_font(styles["Title"], size=24, color=INK, bold=True)
    styles["Title"].paragraph_format.space_after = Pt(4)

    set_style_font(styles["Subtitle"], size=12, color=MUTED, bold=False)
    styles["Subtitle"].paragraph_format.space_after = Pt(18)

    set_style_font(styles["Heading 1"], size=16, color=BLUE, bold=True)
    styles["Heading 1"].paragraph_format.space_before = Pt(18)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)

    set_style_font(styles["Heading 2"], size=13, color=BLUE, bold=True)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)

    set_style_font(styles["Heading 3"], size=12, color=DARK_BLUE, bold=True)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("CoinDCX INR-M Futures Bot | Confidential Pitch Brief")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Prepared for discussion only. Not financial advice. Live trading requires further validation.")
    set_run_font(run, size=8.5, color=MUTED)


def build():
    doc = Document()
    configure_document(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("PROJECT PITCH")
    set_run_font(run, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("CoinDCX INR-M Futures Trading Bot")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("An explainable, risk-controlled automation platform for crypto futures research, paper trading, and future live execution")

    add_label_table(
        doc,
        [
            ("Current stage", "Advanced paper-trading and backtesting system. Not approved for unsupervised live trading yet."),
            ("Market focus", "CoinDCX INR-M crypto futures, built for Indian trading constraints and INR margin workflows."),
            ("Core thesis", "A disciplined automation stack can reduce emotional execution, enforce risk limits, and turn strategy research into testable, auditable trading decisions."),
            ("Prepared date", date.today().strftime("%B %d, %Y")),
        ],
    )

    add_callout(
        doc,
        "One-sentence pitch",
        "This bot is a Python-based CoinDCX futures automation system that combines live market data, explainable hybrid strategy scoring, risk-first position sizing, paper execution, realistic fee/slippage modeling, and a web dashboard into one deployable trading research platform.",
        fill=LIGHT_BLUE_FILL,
    )

    doc.add_heading("Executive Summary", level=1)
    add_para(
        doc,
        "The CoinDCX INR-M Futures Trading Bot is being built as a full-cycle automated trading platform: it ingests market data, evaluates strategy signals, sizes risk, simulates or routes orders, tracks positions, and reports performance through a local dashboard. The project is deliberately staged. Today, the bot is strongest as a research, validation, and paper-trading system. The live-trading layer is intentionally locked until reconciliation, exchange-order monitoring, stronger persistence, and production kill-switches are fully validated.",
    )
    add_para(
        doc,
        "The project is not positioned as a guaranteed-return product. Its value is the engineering system: repeatable strategy testing, disciplined risk controls, auditability, and an architecture that can mature toward live execution without skipping the safety layer.",
    )

    doc.add_heading("The Problem", level=1)
    for item in [
        "Manual futures trading is emotionally noisy: entries, exits, and position sizing are often inconsistent when volatility spikes.",
        "Most retail bot experiments over-focus on signals and under-build the operational layer: slippage, fees, margin checks, account depletion, and restart recovery.",
        "Indian crypto traders who prefer INR-M futures need fee, GST, margin currency, and exchange-specific assumptions modeled directly instead of copied from USDT systems.",
        "Backtesting alone can be misleading if it ignores same-candle ambiguity, fees, slippage, stop slippage, daily loss limits, and realistic execution sequencing.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("The Solution", level=1)
    add_para(
        doc,
        "The bot is designed as a modular automation stack rather than a single monolithic strategy script. Each decision has a path: market data enters the pipeline, indicators and strategy components produce a signal, the risk manager validates exposure, the paper/live broker executes or rejects, and the dashboard records exactly what happened.",
    )
    add_matrix(
        doc,
        ["Layer", "What it does", "Why it matters"],
        [
            ("Market data", "CoinDCX REST and WebSocket data, candles, order book depth, and live feeds.", "Keeps research and paper trading close to exchange behavior."),
            ("Strategy engine", "Hybrid Meta V2, adaptive/weighted hybrid logic, EMA/RSI, Bollinger, visual screen, ATR policy routing.", "Allows strategy ideas to be tested as explainable components, not black boxes."),
            ("Risk manager", "Risk per trade, daily loss guard, leverage cap, margin sufficiency, account depletion checks, exposure limits.", "Prevents the strategy from bypassing capital protection."),
            ("Paper broker", "Simulated fills, maker/taker fees, GST, funding placeholder, order-book VWAP slippage, stop and trailing logic.", "Tests execution behavior without real capital."),
            ("Dashboard", "Backtests, paper trading monitor, equity curve, live candle curve, trades timeline, settings controls.", "Makes the bot inspectable and usable during validation."),
            ("Persistence and logs", "Paper session state, trade history, backtest history, rotating logs.", "Supports auditing, restart recovery, and result review."),
        ],
        [1500, 4100, 3760],
    )

    doc.add_heading("Current Capabilities", level=1)
    add_matrix(
        doc,
        ["Capability", "Status", "Notes"],
        [
            ("CoinDCX API integration", "Built", "Private auth, public data, futures-focused configuration, WebSocket channels."),
            ("Backtesting", "Built", "Supports fees, GST, slippage, ATR exits, intrabar execution, result export/history."),
            ("Paper trading", "Built and active", "Runs with fake capital but live market data; live orders remain locked."),
            ("Dashboard", "Built", "Local web interface for backtests, paper trading, trades, positions, equity, and candles."),
            ("Strategy stack", "Active research", "Hybrid Meta V2 is the main direction; adaptive/weighted hybrids remain available for comparison."),
            ("Risk controls", "Built", "Includes margin sufficiency, equity depletion, leverage guard, daily loss, stop-loss, trailing, profit lock."),
            ("Live trading", "Not ready", "Requires order idempotency, exchange reconciliation, production monitoring, and kill-switch hardening."),
        ],
        [2500, 1800, 5060],
    )

    doc.add_heading("Strategy Approach", level=1)
    add_para(
        doc,
        "The primary strategy direction is Hybrid Meta V2. It blends multiple independent views of the market rather than relying on one indicator. EMA/RSI provides trend and momentum context. Bollinger logic detects reversion or compression behavior. Visual screening evaluates structure, volume, candle quality, and whether the setup looks too weak or too extended. ATR policy routing chooses the exit style per trade.",
    )
    add_matrix(
        doc,
        ["Component", "Role in the decision", "Current treatment"],
        [
            ("EMA/RSI", "Trend and momentum bias.", "Useful especially on cleaner higher-timeframe regimes."),
            ("Bollinger logic", "Mean reversion, compression, and range behavior.", "Kept as a component but no longer assumed to be strong on all long intervals."),
            ("Visual screen", "Candle shape, volume, structure, and obvious chase/weakness filters.", "Used to reduce low-quality entries."),
            ("Open interest", "Optional confirmation when available.", "CoinDCX/Binance proxy availability is inconsistent for some INR-M pairs, so it cannot be mandatory everywhere."),
            ("ATR router", "Selects runner, defensive, or target-style exit behavior.", "Deterministic and metadata-driven so the selected policy is visible per trade."),
            ("Intrabar breakout", "Allows reduced-risk long entries when 1m impulse breaks parent context early.", "Now protected by reduced risk, profit lock, ATR trailing, and stagnation stop instead of blunt time stop."),
        ],
        [2200, 3900, 3260],
    )

    doc.add_heading("Risk Management Philosophy", level=1)
    add_callout(
        doc,
        "Risk-first design",
        "The bot is built so a strategy signal is never enough by itself. Every entry must pass risk validation before execution. This protects against oversized trades, margin insufficiency, account depletion, excessive exposure, and daily loss breaches.",
    )
    for item in [
        "Risk per trade is configurable and can be reduced automatically for weaker or special-case setups.",
        "Position sizing uses initial equity by default, preventing risk from automatically compounding after profits.",
        "Margin sufficiency blocks entries when required margin or planned risk exceeds available equity.",
        "Daily loss guard can pause or force-close behavior when the session breaches the configured limit.",
        "ATR exits, profit lock, breakeven, trailing stops, and stagnation stops are executed by the broker layer, not only displayed in the dashboard.",
        "Live orders are locked by default until production-grade exchange reconciliation is implemented.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Why This Is Different", level=1)
    add_matrix(
        doc,
        ["Differentiator", "Why it is valuable"],
        [
            ("INR-M first", "Fees, GST, margin currency, and CoinDCX pair conventions are built into the system instead of treated as afterthoughts."),
            ("Explainable signals", "Trade metadata stores scores, ATR profile, risk multiplier, exit policy, and rejection reasons."),
            ("Realistic paper execution", "Paper mode now uses live order-book depth to estimate market-order slippage instead of relying only on fixed percentages."),
            ("Signal funnel", "The system can show where candidates are rejected: threshold, visual screen, cooldown, risk, margin, OI, or existing position."),
            ("Safety before live", "The project deliberately blocks live trading until the operational controls are strong enough."),
        ],
        [2600, 6760],
    )

    doc.add_heading("Validation So Far", level=1)
    add_para(
        doc,
        "The bot has gone through several rounds of code review, audit validation, bug fixing, and behavior tuning. Backtesting showed that not all strategies generalize equally across coins or intervals. Early results suggested stronger behavior in selected 1h cases and inconsistent performance in shorter timeframes. More recently, the focus shifted from chasing higher trade counts to measuring why trades are accepted or blocked, improving execution realism, and protecting against repeated loss sequences.",
    )
    add_matrix(
        doc,
        ["Validation area", "What has been checked"],
        [
            ("API/auth", "CoinDCX API key auth was validated after resolving IP whitelist and clock issues."),
            ("Backtest correctness", "Fees, slippage, GST, maker/taker treatment, intrabar execution, and trade history export have been iterated."),
            ("Risk controls", "Margin sufficiency, account depletion, risk base, max daily loss, leverage checks, and exposure checks have been added or strengthened."),
            ("Paper trading", "Multi-pair paper mode, live candle history, equity display, paper session persistence, and order-book slippage were improved."),
            ("Tests", "Focused unit and integration tests cover strategy selection, broker fills, risk behavior, intrabar behavior, dashboard wiring, and market data normalization."),
        ],
        [2500, 6860],
    )

    doc.add_heading("Current Limitations", level=1)
    add_callout(
        doc,
        "Important caveat",
        "This is not live-capital ready yet. It is a serious prototype and paper-trading system. Moving to live trading before reconciliation, order idempotency, exchange-side monitoring, and production kill-switches would be premature.",
        fill="FFF2CC",
    )
    for item in [
        "Open-interest data is not reliable for every CoinDCX INR-M pair, so OI must remain optional or verified per market.",
        "Paper slippage can approximate live slippage from visible order-book depth, but it cannot perfectly model latency, queue priority, hidden liquidity, or matching-engine behavior.",
        "Backtests can still be regime-sensitive. A strategy working on one interval or coin does not prove broad robustness.",
        "Live trading still needs full order lifecycle tracking: submitted, acknowledged, partially filled, filled, canceled, rejected, and reconciled.",
        "Production deployment needs heartbeat alerts, crash recovery, log monitoring, and a human-controlled emergency shutdown process.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Roadmap To Live Trading", level=1)
    for item in [
        "Complete paper-trading observation across selected pairs and intervals, with every trade logged and reviewed.",
        "Lock a conservative strategy profile based on out-of-sample behavior, not a single profitable backtest window.",
        "Add live exchange reconciliation: compare local positions, open orders, fills, and balances against CoinDCX at intervals.",
        "Add idempotent order IDs and retry-safe execution so a network retry cannot accidentally double-enter.",
        "Add Telegram alerts for entries, exits, errors, daily summary, drawdown warnings, and kill-switch triggers.",
        "Run a tiny-capital supervised pilot only after paper behavior, reconciliation, and kill switches are stable.",
    ]:
        add_number(doc, item)

    doc.add_heading("Suggested 60-Second Pitch", level=1)
    add_callout(
        doc,
        "Pitch script",
        "I am building an INR-M CoinDCX futures automation platform, not just a single trading strategy. The system watches live markets, evaluates hybrid strategy signals, checks risk and margin before every entry, simulates execution with realistic fees and order-book slippage, and records every trade in a dashboard. The near-term goal is disciplined paper validation and strategy research. The long-term goal is a production-ready live trading engine with reconciliation, kill-switches, and fully auditable decision logs. The edge I am trying to build is not blind prediction. It is repeatable execution, risk discipline, and a system that can tell us exactly why it traded or why it stayed out.",
        fill=LIGHT_BLUE_FILL,
    )

    doc.add_heading("What I Am Looking For", level=1)
    for item in [
        "Technical review from someone experienced in exchange execution, trading-system reliability, or crypto derivatives.",
        "Help evaluating whether Hybrid Meta V2 has enough out-of-sample robustness to justify a supervised live pilot.",
        "Guidance on CoinDCX-specific execution details: order types, leverage constraints, tick/step sizes, and WebSocket reliability.",
        "Optional small-capital pilot support only after paper trading proves stable and all live safeguards are complete.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Final Positioning", level=1)
    add_para(
        doc,
        "This project should be pitched as a serious trading automation platform under validation, not a finished yield product. The strongest claim is that it has the architecture, controls, and auditability needed to develop strategies responsibly. The next milestone is not simply higher profit. It is proving that the bot can survive real market behavior, avoid destructive execution mistakes, and make decisions that remain explainable after the fact.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
