from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "coindcx_inr_margin_position_sizing_explainer.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(18, 24, 38)
MUTED = RGBColor(83, 96, 117)
FILL = "F2F4F7"
CALLOUT = "EAF4FF"
WARN = "FFF4E5"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_grid(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_text(cell, text: str, *, bold: bool = False, color: RGBColor | None = None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    set_table_grid(table, widths)
    for idx, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color=INK)
        set_cell_shading(table.rows[0].cells[idx], FILL)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            set_cell_text(cells[idx], value)
    doc.add_paragraph()


def add_callout(doc: Document, title: str, body: str, fill: str = CALLOUT) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_grid(table, [9360])
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    r.font.color.rgb = DARK_BLUE
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(body)
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.color.rgb = INK
    doc.add_paragraph()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = BLUE if level <= 2 else DARK_BLUE


def add_body(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = INK


def add_number(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = INK


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(4)
    r = title.add_run("CoinDCX INR-M Futures Bot")
    r.font.name = "Calibri"
    r.font.size = Pt(24)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    sr = subtitle.add_run("Position sizing, margin, notional, quantity, fees, and risk explained using your current dashboard setup")
    sr.font.name = "Calibri"
    sr.font.size = Pt(11)
    sr.font.color.rgb = MUTED

    add_callout(
        doc,
        "The one-line answer",
        "The bot does not invest your full Rs. 10,000 in one trade. With your current setup, it first decides the maximum planned loss per trade, then calculates quantity, then checks whether the required margin is allowed.",
    )

    add_heading(doc, "1. Your Current Dashboard Setup", 1)
    add_table(
        doc,
        ["Dashboard field", "Current value", "What it means"],
        [
            ["Starting Equity", "Rs. 10,000", "The paper account balance used as the capital base."],
            ["Risk / Trade", "3%", "Normal trades may risk up to Rs. 300 if the stop-loss is hit."],
            ["Leverage", "5x", "The position value can be up to 5 times the margin used."],
            ["Max Daily Loss", "10%", "Bot should stop new entries after roughly Rs. 1,000 daily loss."],
            ["Max Open Positions", "3", "Up to 3 open positions across the watchlist."],
            ["Max Margin", "50%", "Maximum total margin usage is Rs. 5,000 from Rs. 10,000 equity."],
            ["Multi-Pair", "On", "BSB and AXL can both be traded."],
            ["Pyramiding", "Off", "Bot should not keep adding to the same open pair."],
            ["Strategy Interval", "5m", "Strategy decision context is based on 5-minute candles."],
            ["Intrabar Execution", "1m", "Entries and exits can execute on 1-minute candles inside the 5-minute setup."],
        ],
        [2100, 1600, 5660],
    )

    add_heading(doc, "2. The Four Money Numbers", 1)
    add_body(doc, "Most confusion comes from mixing these four different numbers. They are related, but they are not the same.")
    add_table(
        doc,
        ["Term", "Simple meaning", "Example"],
        [
            ["Equity", "Total paper account value.", "Rs. 10,000"],
            ["Risk", "Maximum planned loss if stop-loss hits.", "3% of Rs. 10,000 = Rs. 300"],
            ["Position Notional", "Total trade value controlled by the position.", "About Rs. 3,037 in the BSB example"],
            ["Margin", "Amount blocked to hold the leveraged position.", "Rs. 3,037 / 5 = about Rs. 607"],
        ],
        [1700, 4300, 3360],
    )

    add_callout(
        doc,
        "Important distinction",
        "Position notional is the size of the trade. Margin is the amount blocked. Risk is the amount the bot expects to lose if the stop-loss is hit. Equity is your full account value.",
        fill=WARN,
    )

    add_heading(doc, "3. Why Quantity Comes Before Margin", 1)
    add_body(doc, "The bot is designed to size trades from risk first. This is safer than choosing quantity from available margin because it prevents one bad stop-loss from taking an oversized loss.")
    add_number(doc, "Calculate maximum planned loss: Rs. 10,000 x 3% = Rs. 300.")
    add_number(doc, "Look at the entry price and stop-loss price.")
    add_number(doc, "Calculate the loss per coin if stop-loss is hit.")
    add_number(doc, "Choose quantity so the total stop-loss loss is about Rs. 300.")
    add_number(doc, "Calculate position notional from that quantity.")
    add_number(doc, "Calculate required margin from notional / leverage.")
    add_number(doc, "Reject the trade if margin, exposure, daily loss, or open-position rules are violated.")

    add_heading(doc, "4. Worked BSB Example", 1)
    add_body(doc, "Assume a BSB short trade similar to the dashboard numbers. Prices are USDT-quoted because CoinDCX INR-M futures still show the chart price in USDT. The account, risk, margin, PnL, and fees are INR.")
    add_table(
        doc,
        ["Calculation", "Formula", "Result"],
        [
            ["Account equity", "Given", "Rs. 10,000"],
            ["Risk allowed", "Rs. 10,000 x 3%", "Rs. 300"],
            ["Entry price", "Given by signal", "0.632094 USDT"],
            ["Stop price", "Given by signal", "0.694538 USDT"],
            ["Stop distance", "0.694538 - 0.632094", "0.062444 USDT"],
            ["USDT to INR rate", "Config value", "About Rs. 98 per USDT"],
            ["Risk per BSB", "0.062444 x 98", "About Rs. 6.12"],
            ["Quantity", "Rs. 300 / Rs. 6.12", "About 49.02 BSB"],
            ["Position notional", "49.02 x 0.632094 x 98", "About Rs. 3,036.77"],
            ["Required margin", "Rs. 3,036.77 / 5", "About Rs. 607.35"],
        ],
        [2100, 3450, 3810],
    )

    add_callout(
        doc,
        "What changed after the bug fix",
        "Before the fix, quantity could look like around 5,000 BSB because the bot treated a USDT price movement as if it were already INR. After the fix, the same setup is around 49 BSB, with PnL and margin still shown in INR.",
    )

    add_heading(doc, "5. Why AXL Quantity Can Look Bigger Than BSB", 1)
    add_body(doc, "A smaller coin price naturally needs more coins to reach the same INR notional. That does not mean the trade is larger or riskier.")
    add_table(
        doc,
        ["Pair", "Approx price", "Approx quantity for Rs. 3,000 notional", "Meaning"],
        [
            ["BSB", "About Rs. 62", "About 48 to 50 BSB", "Higher coin price, smaller quantity."],
            ["AXL", "About Rs. 6", "About 500 AXL", "Lower coin price, larger quantity."],
        ],
        [1200, 1800, 3000, 3360],
    )

    add_heading(doc, "6. What Leverage And Max Margin Actually Do", 1)
    add_body(doc, "Leverage does not decide how much you are willing to lose. The stop-loss risk decides that. Leverage mainly decides how much margin is needed to carry the position.")
    add_table(
        doc,
        ["Control", "Your value", "Practical effect"],
        [
            ["Risk / Trade", "3%", "A normal trade should risk about Rs. 300."],
            ["Leverage", "5x", "A Rs. 3,000 position needs about Rs. 600 margin."],
            ["Max Margin", "50%", "Total margin across open trades should stay under about Rs. 5,000."],
            ["Max Open Positions", "3", "The bot can split risk across several pairs, but each trade still needs risk and margin approval."],
        ],
        [2100, 1600, 5660],
    )

    add_heading(doc, "7. Why You Might See Less Than 3% Risk", 1)
    add_body(doc, "The dashboard says 3% risk per trade, but the strategy can deliberately reduce risk for weaker or more aggressive entries.")
    add_bullet(doc, "A strong setup may use the full 3% risk, about Rs. 300.")
    add_bullet(doc, "A weaker B setup may use a smaller multiplier, such as 50%, so risk becomes about Rs. 150.")
    add_bullet(doc, "A breakout or defensive setup may use even less if the strategy labels it as risky.")
    add_bullet(doc, "This is intentional: the strategy is allowed to reduce risk, but it should not exceed the configured risk cap.")

    add_heading(doc, "8. Dashboard Reading Checklist", 1)
    add_bullet(doc, "Quantity is coin/contracts. It is not rupees.")
    add_bullet(doc, "Notional is the total INR value of the position.")
    add_bullet(doc, "Required margin is notional divided by leverage.")
    add_bullet(doc, "Risk used is the planned INR loss if stop-loss hits.")
    add_bullet(doc, "PnL, fees, equity, notional, and margin should all be read as INR.")
    add_bullet(doc, "Chart prices and entry/exit prices remain USDT-quoted because that is how CoinDCX shows these futures markets.")

    add_heading(doc, "9. What To Do After This Fix", 1)
    add_callout(
        doc,
        "Reset the paper session once",
        "Old open paper positions were created using the old quantity math. Resetting the paper session makes the next trades clean and comparable under the corrected INR model.",
        fill=WARN,
    )
    add_body(doc, "After reset, check that BSB quantity is no longer thousands of coins for a small Rs. 10,000 account trade. It should be closer to the size implied by risk, stop distance, and the INR conversion rate.")

    doc.add_section(WD_SECTION.CONTINUOUS)
    footer = doc.sections[-1].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = footer.add_run("CoinDCX INR-M futures bot explainer")
    fr.font.name = "Calibri"
    fr.font.size = Pt(9)
    fr.font.color.rgb = MUTED

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
